"""Build the editable one-slide team pitch, including PowerPoint speaker notes.

From the repository root:
  ./.venv/Scripts/python.exe -m pip install python-pptx==1.0.2
  ./.venv/Scripts/python.exe scripts/pitch/build_one_slide.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs/pitch/swisstip-one-slide.pptx"

BACKGROUND = "F8F9F5"
INK = "173E33"
MUTED = "586C62"
ACCENT = "DDEBAD"
DIVIDER = "D8E0D7"

NOTES = """SUGGESTED TALK TRACK (ABOUT 2 MINUTES)

I recommend choosing Swisscom's Swiss Grounding MCP challenge because it combines a relatable Swiss problem, substantial engineering work and a credible route beyond the hackathon.

Imagine asking: "How to get Aufenthaltsbewilligung in Zurich?" A useful answer depends on the right authority and the person's actual situation. Our assistant would clarify the missing details, use TIP to retrieve scoped official evidence, and explain what that evidence supports. The user can inspect the original source and see any remaining gaps.

The slide makes our case against each jury criterion. For technical functionality, we can show working source ingestion, retrieval and MCP components. For skillful use of AI, models contribute to extraction, retrieval and ranking. For user experience, our goal is to turn complex Swiss processes into simple, guided answers. For creativity, we can let people ask in their language and discover the relevant Swiss context, with a memorable reveal of the original source. These user and multilingual benefits are the experience we are completing, not a claim of fully validated coverage today.

The commercial argument is equally concrete. Swisscom has stated that it wants a testable prototype with potential later integration into myAI, and the same evidence service could support other assistants and applications. That gives us an adoption hypothesis and a reason to build something reusable.

Our next priority is to finish one polished Zurich arrival journey with reviewed evidence and repeatable setup. This gives us a clear way to demonstrate all five jury criteria while bringing practical AI, provenance and integration lessons back to UBS.

JURY MAPPING

Technical functionality: demonstrate several connected steps from official source acquisition and saved snapshots through preparation, scoped retrieval and the real MCP response. The individual components are implemented; a complete reviewed publication pipeline remains unfinished.

Skillful use of AI: models propose knowledge during extraction and assist multilingual retrieval and ranking. The calling assistant handles conversation. Human review and deterministic scope checks limit what model output can establish. MCP is the interface for AI tool use; it is not, by itself, the AI contribution.

User experience: show a brief conversation, only necessary clarifications, readable evidence, a clickable source and clear loading or failure states. The existing knowledge studio is operator UX. The complete end-user assistant experience still needs qualification.

Uniqueness, creativity and fun: make the "show me the source" moment memorable. Use mixed English/German wording, then ask for a fee absent from the selected evidence. Explain that the selected evidence is insufficient; do not claim the authority has no such information. The distinctive proposition is the combination of Swiss relevance, reusable evidence and visible limits. MCP, RAG and citations alone are not exclusive inventions.

Potential and market impact: myAI is a possible adoption path stated by the challenge, not a confirmed integration or customer commitment. Other possible users include public-service and relocation applications and enterprise assistant teams. Maintained knowledge and integration are commercial hypotheses; demand and willingness to pay have not been established.

JURY PREFERENCES AND DELIVERY FOCUS

Reproducibility: existing synthetic-fixture setup and automated tests provide a baseline. Complete a clean-checkout rehearsal and a permitted corpus or rebuild procedure for the real demonstration. Label any recorded replay. A Git-ignored local pilot archive is not a portable submission.

AI paradigm: demonstrate how an agent discovers and consumes a reusable evidence service, with a clear split between conversational interpretation and evidence retrieval.

Challenge fit: focus on authoritative Swiss sources, relevant evidence and efficient reuse through MCP. Verify the sponsor's evaluation harness when available. The proposed TIP request schema is our design choice.

Keep the corpus narrow while completing accepted requirements, including reviewed five-language metadata. A narrow demo does not silently remove that P0 requirement. Nationwide coverage, a marketplace and production enterprise deployment remain future work.

CURRENT EVIDENCE AND LIMITS - REVIEWED 2026-09-09

The SEM/Zurich pilot retrieved real saved excerpts and withheld evidence for selected unsupported queries. Its catalog IDs, coverage and approval references remain experimental. It has zero published facts and rules, so it does not prove supported individual legal guidance. The final Zurich observations were collected across runs, including an operational failure followed by a passing rerun. Guided caller checks required corrections. Five-language draft projections are not independent language-quality validation.

The review reran 67 passing runtime tests and four passing MCP tests. Six database integration tests were skipped because their test connection was not configured. These are implementation checks, not legal or live-model quality proof.

Avoid claims of nationwide coverage, guaranteed accuracy, five-language quality already proven, production readiness, confirmed myAI adoption or a guaranteed win. The jury's criteria are qualitative and do not define fixed weights.

SOURCES AND BACKGROUND

Official challenge listing: https://ai-weeks.ch/2026/challenges
The organizer's indexed listing confirms the grounding problem, sponsor-testable prototype and possible later myAI integration. Retrieved during the documentation review on 2026-09-09; detailed harness requirements were not independently revalidated.

Jury criteria: supplied by the user in this project discussion.
docs/pitch/one-page-pitch.md
docs/pitch/project-review.md
docs/pilots/2026-09-08-retrieval-pilot.md
apps/mcp-server/README.md
apps/admin-console/README.md
docs/product/product-functional-specification.md
"""


def color(value):
    return RGBColor.from_string(value)


def rectangle(slide, x, y, width, height, fill):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(width), Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    # Override the template's default shadow for a flat, simple slide.
    shape._element.spPr.get_or_add_effectLst()
    return shape


def text(slide, value, x, y, width, height, size, *, bold=False, fill=INK):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height),
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(fill)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.12
    return shape


def build():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "SwissTIP - Why choose Swiss Grounding MCP"
    presentation.core_properties.subject = "One-slide challenge-selection pitch with speaker notes"
    presentation.core_properties.author = "SwissTIP team"
    presentation.core_properties.keywords = "SwissTIP, Swiss Grounding MCP, hackathon, jury criteria"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BACKGROUND)

    rectangle(slide, 0.64, 0.48, 0.13, 0.18, INK)
    text(slide, "SWISSTIP", 0.87, 0.435, 2.0, 0.28, 13, bold=True)
    text(slide, "TEAM CHALLENGE CHOICE", 9.05, 0.45, 3.6, 0.26, 11, fill=MUTED)

    text(slide, "Why this challenge can win.",
         0.64, 1.04, 12.05, 0.79, 40, bold=True)
    text(slide, "SwissTIP | Swisscom Swiss Grounding MCP",
         0.68, 1.98, 11.9, 0.4, 21, fill=MUTED)

    # Five explicit jury criteria plus the published preferences, in a quiet grid.
    for x in (4.56, 8.62):
        rectangle(slide, x, 2.95, 0.014, 3.49, DIVIDER)
    rectangle(slide, 0.68, 4.62, 11.96, 0.014, DIVIDER)
    rectangle(slide, 4.84, 4.86, 3.5, 1.59, ACCENT)

    criteria = [
        (0.68, 2.98, "TECHNICAL FUNCTIONALITY",
         "Working source ingestion,\nretrieval and MCP", 20),
        (4.99, 2.98, "SKILLFUL USE OF AI",
         "AI extracts, retrieves\nand ranks evidence", 20),
        (9.01, 2.98, "USER EXPERIENCE",
         "Complex Swiss processes.\nSimple, guided answers.", 20),
        (0.68, 5.03, "UNIQUENESS, CREATIVITY & FUN",
         "Ask in your language.\nGet the Swiss context.", 20),
        (4.99, 5.03, "POTENTIAL / MARKET IMPACT",
         "Potential myAI integration\nReuse across assistants", 19),
        (9.01, 5.03, "JURY PREFERENCES",
         "Reproducible code\nAI-native design\nDirect challenge fit", 19),
    ]
    for x, y, label, body, size in criteria:
        text(slide, label, x, y, 3.63, 0.26, 10.6, bold=True, fill=MUTED)
        text(slide, body, x, y + 0.55, 3.63, 1.05, size, bold=True)

    slide.notes_slide.notes_text_frame.text = NOTES.strip()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
